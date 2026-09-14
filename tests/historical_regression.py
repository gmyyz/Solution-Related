"""Opt-in local historical regression; all writes stay in a temporary directory.

Usage: python tests/historical_regression.py [2603 2606 2607 2608]
No production data or amounts are written to this repository.
"""
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tarfile
import tempfile

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'工作区域'
RUNNER=r'''
import sys,json,hashlib
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0,str(Path('工作区域').resolve()))
from payroll_tool.options import RunOptions,build_batch_layout_from_options
from payroll_tool.pipeline import execute_payroll_run
from payroll_tool.core import _find_raw_files
month=int(sys.argv[1])
options=RunOptions(('耐数电子','耐数信息'),tuple('A'+str(i) for i in range(1,11)),2026,month)
base=str(Path('工作区域').resolve());layout=build_batch_layout_from_options(base,options)
try:
    paths=_find_raw_files(layout.raw_dir,2026,month-1)
    result=execute_payroll_run(paths[0],base,layout.mapping_path,layout.bank_dir,lambda *args:None,run_options=options)
    vouchers={}
    for category,path in result['voucher_paths'].items():
        with open(path,'rb') as f:data=f.read()
        import io
        wb=load_workbook(io.BytesIO(data),data_only=True)
        try:
            rows=list(wb.active.iter_rows(min_row=2,max_col=17,values_only=True))
            rows=[r for r in rows if any(v is not None for v in r)]
            vouchers[category]={'rows':len(rows),'sha256':hashlib.sha256(json.dumps(rows,default=str,ensure_ascii=True).encode()).hexdigest()}
        finally:wb.close()
    print(json.dumps({'status':'generated','vouchers':vouchers},ensure_ascii=True))
except Exception as e:
    print(json.dumps({'status':'blocked','error':type(e).__name__+': '+str(e)},ensure_ascii=True))
'''


def main():
    codes=sys.argv[1:] or ['2603','2606','2607','2608']
    results=[]
    tracked={}
    def source_copy(source,target):
        if not source.is_file():return
        with open(source,'rb') as stream:tracked[str(source)]=hashlib.file_digest(stream,'sha256').hexdigest()
        target.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(source,target)
    with tempfile.TemporaryDirectory(prefix='payroll-history-') as temp:
        for code in codes:
            batch=next((WORK/'02_月度输入').glob('*('+code+')'))
            comparison={'batch':code}
            for version in ['baseline','current']:
                root=Path(temp)/code/version;root.mkdir(parents=True)
                if version=='baseline':
                    archive=subprocess.check_output(['git','archive','19e373e','工作区域/payroll_tool'],cwd=ROOT)
                    with tarfile.open(fileobj=io.BytesIO(archive)) as tar:tar.extractall(root,filter='data')
                else:
                    shutil.copytree(WORK/'payroll_tool',root/'工作区域/payroll_tool',ignore=shutil.ignore_patterns('__pycache__'))
                for source in batch.rglob('*.xlsx'):source_copy(source,root/'工作区域'/source.relative_to(WORK))
                for relative in ['01_基础资料/01_映射与规则/Mapping表.xlsx','01_基础资料/02_工时数据/工时数据.xlsx','Mapping表.xlsx']:
                    source_copy(WORK/relative,root/'工作区域'/relative)
                source_copy(ROOT/'模版性文件/2060总账凭证导入-实际202602薪酬.XLS',root/'模版性文件/2060总账凭证导入-实际202602薪酬.XLS')
                try:
                    completed=subprocess.run([sys.executable,'-c',RUNNER,str(int(code[-2:]))],cwd=root,capture_output=True,timeout=90)
                    comparison[version]=json.loads(completed.stdout.decode('utf8').strip().splitlines()[-1])
                except subprocess.TimeoutExpired:
                    comparison[version]={'status':'timeout','limit_seconds':90}
                except Exception as exc:
                    comparison[version]={'status':'failed','error':str(exc)}
                print(json.dumps({'batch':code,'version':version,'status':comparison[version]['status'],
                                  'error':comparison[version].get('error')},ensure_ascii=True),flush=True)
            if all(comparison[v].get('status')=='generated' for v in ['baseline','current']):
                comparison['equal']=comparison['baseline']['vouchers']==comparison['current']['vouchers']
            else:comparison['equal']=None
            results.append(comparison)
    unchanged=True
    for path,expected in tracked.items():
        with open(path,'rb') as stream:unchanged &= hashlib.file_digest(stream,'sha256').hexdigest()==expected
    print(json.dumps({'source_files_unchanged':unchanged,'results':[
        {'batch':r['batch'],'equal':r['equal'],'baseline':r['baseline']['status'],
         'current':r['current']['status'],'voucher_files':len(r['current'].get('vouchers',{}))}
        for r in results]},ensure_ascii=True),flush=True)


if __name__=='__main__':main()
