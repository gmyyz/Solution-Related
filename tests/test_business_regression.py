"""Small, explicit financial fixtures; independent expected amounts."""
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import sys
import tempfile
import unittest
from openpyxl import Workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / '工作区域'))
from payroll_tool import vouchers as v
from payroll_tool.domain import _allocate_amount_by_weights, _should_skip_one_time_june_2026_accrual
from payroll_tool.contexts import _load_co_workorder_context, _load_shared_expense_context
from payroll_tool.validation import _validate_voucher_groups, _validate_cross_group_accounts


def payroll(company='耐数电子', code='2050', dept='研发部', header_row=2):
    wb=Workbook(); ws=wb.active
    if header_row==2: ws.append(['工资所属月份'])
    ws.append(['公司','部门','项目','类别']+['']*16)
    ws.append([company,dept,'项目A','薪资',100,10,0,0,0,2,1,0,0,2,10,0,87,0,code+'1001','9201001'])
    ws.append(['总计','','','']+[ws.cell(header_row+1,c).value for c in range(5,19)])
    return wb,ws,header_row+2


def bank(code):
    def rec(amount,name,usage):
        return dict(outgoing_amt=Decimal(str(amount)), payment_name=name,usage=usage,
                    bank_subject='1002010001',trans_date=datetime(2026,3,20),row_idx=2,bank_account='a')
    return {(code,2026,3):[rec(87,'员工','代发工资'),rec(10,'国家金库北京市分库','税'),
                          rec(11,'国家金库北京市分库','社保'),rec(4,'北京住房公积金管理中心','公积金')]}


def timesheet(company='耐数电子',dept='研发部',month=3):
    return {'allocation_index':{(2026,month,company,dept):{'9201001':Decimal('1')}},
            'project_index':{(company,'项目A'):{'order':'9201001','types':{'研发'}}},
            'kok3_index':{},'extra_lookup':{}}


def expense(company='耐数电子',dept='研发部',cc='20501001',amount='10',order='old'):
    return {'rows_by_company':{company:[dict(row_idx=2,gl_account='6601010001',dept=dept,
             cost_center=cc,allocation_cost_center=cc,credit_order=order,amount=Decimal(amount))]}}


class BusinessRegressionTests(unittest.TestCase):
    def balanced(self,rows,total):
        _validate_voucher_groups(rows)
        for key in (40,50):
            self.assertEqual(sum((Decimal(r['M']) for r in rows if r['G']==key),Decimal('0')),Decimal(str(total)))
        self.assertTrue(all(isinstance(r['M'],str) for r in rows))

    def test_a1_through_a8_both_companies_and_header_formats(self):
        for company,code in [('耐数电子','2050'),('耐数信息','2060')]:
            for header in [1,2]:
                with self.subTest(company=company,header=header):
                    wb,ws,total=payroll(company,code,header_row=header)
                    try:
                        b=bank(code); t=timesheet(company)
                        rows=[v._build_a1_rows(company,code,ws,total,b,2026,2),
                              v._build_a2_rows(company,code,b,2026,2),
                              v._build_a3_rows(company,code,ws,b,2026,2,t),
                              v._build_a4_rows(company,code,ws,total,2026,2,t),
                              v._build_a5_rows(company,code,ws,total,2026,2),
                              v._build_a6_rows(company,code,ws,2026,2,t)]
                        for built,expected in zip(rows,[11,4,97,112,112,112]): self.balanced(built,expected)
                        self.assertEqual(_validate_cross_group_accounts(sum(rows[:4],[])),[])
                        for built in rows[:3]:
                            self.assertTrue(all(r['C']=='KZ' for r in built))
                            self.assertTrue(all(r['P']=='202' for r in built if r['G']==50))
                        bonus={'ratio_map':{3:Decimal('1')},'rows_by_company':{company:[
                            dict(dept='研发部',cost_center=code+'1001',monthly_amounts={1:Decimal(10),2:Decimal(10),3:Decimal(10)})]}}
                        self.balanced(v._build_a7_rows(company,code,2026,2,bonus),30)
                        self.balanced(v._build_a8_rows(company,code,2026,2,bonus,t),30)
                    finally: wb.close()

    def test_a10_netting_and_distinct_credit_orders(self):
        context=expense(); first=context['rows_by_company']['耐数电子'][0]
        context['rows_by_company']['耐数电子'] += [dict(first,amount=Decimal('-2')),dict(first,credit_order='',amount=Decimal('3'))]
        rows=v._build_a10_rows('耐数电子','2050',2026,2,context,timesheet())
        self.balanced(rows,11)
        self.assertEqual({r['K']:r['M'] for r in rows if r['G']==50},{'old':'8.00','':'3.00'})
        self.assertTrue(all(r['C']=='SA' for r in rows))

    def test_a10_zero_negative_and_no_hours(self):
        self.assertEqual(v._build_a10_rows('耐数电子','2050',2026,2,expense(amount='0'),timesheet()),[])
        for context,t in [(expense(amount='-1'),timesheet()),(expense(),timesheet(month=4))]:
            with self.assertRaises(ValueError):v._build_a10_rows('耐数电子','2050',2026,2,context,t)

    def test_a10_hld_exclusion_and_cost_center_translation(self):
        context=expense(dept='HLD硬件逻辑部',cc='20502050')
        context['rows_by_company']['耐数电子'][0]['allocation_cost_center']='20502060'
        t=timesheet(dept='HLD硬件逻辑部'); t['allocation_index'][(2026,3,'耐数电子','HLD硬件逻辑部')]['9201856']=Decimal(999)
        rows=v._build_a10_rows('耐数电子','2050',2026,2,context,t)
        self.balanced(rows,10)
        self.assertEqual({r['J'] for r in rows if r['G']==40},{'20502060'})
        self.assertEqual({r['J'] for r in rows if r['G']==50},{'20502050'})
        self.assertNotIn('9201856',{r['K'] for r in rows if r['G']==40})

    def test_a10_information_company_fallback(self):
        rows=v._build_a10_rows('耐数信息','2060',2026,2,expense('耐数信息',cc='20602020'),timesheet('耐数信息','其他部门'))
        self.balanced(rows,10)

    def test_wls_exceptions_extend_through_september(self):
        for month in [6,7,8,9]:
            self.assertEqual(_should_skip_one_time_june_2026_accrual('耐数信息',2026,month,'WLS无线解决方案'),month in (6,7))
            rows=v._build_a10_rows('耐数信息','2060',2026,month-1,
                                  expense('耐数信息','WLS无线解决方案','20601001'),
                                  timesheet('耐数信息','WLS无线解决方案',month))
            self.assertEqual(rows,[])

    def test_september_information_ats_accrual_is_skipped(self):
        self.assertTrue(
            _should_skip_one_time_june_2026_accrual(
                '耐数信息', 2026, 9, 'ATS测试解决方案'
            )
        )
        self.assertFalse(
            _should_skip_one_time_june_2026_accrual(
                '耐数信息', 2026, 10, 'ATS测试解决方案'
            )
        )
        self.assertFalse(
            _should_skip_one_time_june_2026_accrual(
                '耐数电子', 2026, 9, 'ATS测试解决方案'
            )
        )

        wb, ws, total = payroll('耐数信息', '2060', 'ATS测试解决方案')
        try:
            self.assertEqual(v._build_a5_rows('耐数信息', '2060', ws, total, 2026, 8), [])
            self.assertEqual(
                v._build_a6_rows('耐数信息', '2060', ws, 2026, 8, timesheet('耐数信息')),
                [],
            )
        finally:
            wb.close()

    def test_a9_loader_mapping_and_marker(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'co.xlsx'; wb=Workbook(); ws=wb.active; ws.title='CO工单分摊'
            ws.append(['订单','','','','借方','','实际']);ws.append(['9201001','','','',1,'',8]);ws.append(['总计','','','',1,'',8])
            ws=wb.create_sheet('待分摊费用'); ws.append(['科目','','','金额']);ws.append(['6601010003','','',5]);ws.append(['other','','',3])
            wb.save(path);wb.close()
            context=_load_co_workorder_context(path)
            rows=v._build_a9_rows('2050',2026,2,context); self.balanced(rows,8)
            self.assertEqual({r['I'] for r in rows if r['G']==40},{'5001010011','5001010012'})
            self.assertTrue(all(r['O']=='x' and r['K']=='9201856' for r in rows if r['G']==50))
            path.unlink()

    def test_a9_negative_expense_uses_absolute_amount_and_swaps_sides(self):
        context = {
            'allocation_rows': [
                {'row_idx': 2, 'order': '9201001', 'ratio_base': Decimal('1'), 'total_actual': Decimal('1')}
            ],
            'total_ratio_base': Decimal('1'),
            'summary_total': Decimal('1'),
            'expense_rows': [
                {
                    'row_idx': 2,
                    'source_account': '6601010003',
                    'debit_gl': '5001010011',
                    'amount': Decimal('-15.10'),
                }
            ],
        }
        rows = v._build_a9_rows('2050', 2026, 8, context)
        self.balanced(rows, Decimal('15.10'))
        self.assertEqual(
            [(r['G'], r['I'], r['J'], r['K'], r['M'], r['O']) for r in rows],
            [
                (40, '6601010003', '20502020', '9201856', '15.10', ''),
                (50, '5001010011', '', '9201001', '15.10', 'x'),
            ],
        )

    def test_small_allocation_and_negative_guard(self):
        self.assertEqual(_allocate_amount_by_weights('.01',[('a',1),('b',1),('c',1)]),{'c':Decimal('.01')})
        self.assertEqual(_allocate_amount_by_weights('1',[('a',0)]),{})
        with self.assertRaises(ValueError):_allocate_amount_by_weights('-1',[('a',1)])

    def test_labor_and_severance_only_affect_actual_payment(self):
        for company,code in [('耐数电子','2050'),('耐数信息','2060')]:
            wb,ws,total=payroll(company,code)
            try:
                ws.insert_rows(total,2)
                for offset,kind,amount,tax,net in [(0,'劳务费',20,2,18),(1,'离职补偿金',30,0,30)]:
                    row=[company,'研发部','项目A',kind,amount]+[0]*9+[tax,0,net,0,code+'1001','9201001']
                    for col,value in enumerate(row,1):ws.cell(total+offset,col,value)
                total+=2
                from payroll_tool.workbook_utils import _recalculate_total_row
                _recalculate_total_row(ws,total)
                b=bank(code); b[(code,2026,3)][0]['outgoing_amt']=Decimal(135)
                b[(code,2026,3)][1]['outgoing_amt']=Decimal(12)
                t=timesheet(company)
                t['project_index'][(company,'研发部日常工作')]={'order':'9201999','types':{'日常'}}
                rows=v._build_a3_rows(company,code,ws,b,2026,2,t)
                self.balanced(rows,147)
                self.assertEqual([(r['M'],r['K']) for r in rows if r['I']=='6601990001'],[('30.00','9201999')])
                self.assertEqual([r['M'] for r in rows if r['I']=='6601110001'],['20.00'])
                self.assertEqual([r['M'] for r in rows if r['I']=='2221070000'],['10.00'])
                self.balanced(v._build_a4_rows(company,code,ws,total,2026,2,t),112)
                self.balanced(v._build_a5_rows(company,code,ws,total,2026,2),112)
                self.balanced(v._build_a6_rows(company,code,ws,2026,2,t),112)
            finally:wb.close()

    def test_september_a3_posts_embedded_bonus_as_separate_rows(self):
        wb,ws,total=payroll()
        try:
            records=bank('2050')[('2050',2026,3)]
            records[0]['outgoing_amt']=Decimal('607920.20')
            september_bank={('2050',2026,9):records}
            rows=v._build_a3_rows('耐数电子','2050',ws,september_bank,2026,8,timesheet())
            self.balanced(rows,Decimal('607930.20'))
            bonus_rows=[r for r in rows if r['N']=='支付2025年9月部分年终奖']
            self.assertEqual(
                [(r['G'],r['I'],r['M']) for r in bonus_rows],
                [(40,'2211010001','607833.20'),(50,'1002010001','607833.20')],
            )
            self.assertEqual(bonus_rows[1]['P'],'202')
        finally:
            wb.close()

    def test_bonus_months_ratio_and_wls_override(self):
        company='耐数信息'; bonus={'ratio_map':{m:Decimal('0.5') for m in range(1,13)},
            'rows_by_company':{company:[dict(dept='研发部',cost_center='20601001',
                                      monthly_amounts={m:Decimal(10) for m in range(1,13)})]}}
        for month,expected in [(1,5),(2,5),(3,15),(6,30),(9,45),(12,60)]:
            if month==1:y,m=2025,12
            else:y,m=2026,month-1
            self.balanced(v._build_a7_rows(company,'2060',y,m,bonus),expected)
        bonus['rows_by_company'][company][0]['dept']='WLS无线解决方案'
        self.assertEqual(v._build_a7_rows(company,'2060',2026,6,bonus),[])
        self.balanced(v._build_a7_rows(company,'2060',2026,7,bonus),5)

    def test_osd_a4_self_research_uses_daily_order(self):
        from payroll_tool.contexts import _resolve_a4_internal_order
        t=timesheet(dept='OSD运营支持部')
        t['project_index'][('耐数电子','项目A')]['types']={'自研'}
        t['project_index'][('耐数电子','OSD运营支持部日常工作')]={'order':'9201999'}
        self.assertEqual(_resolve_a4_internal_order(t,'耐数电子','OSD运营支持部','项目A','9201001'),'9201999')

    def test_osd_is_not_allocated(self):
        wb,ws,total=payroll(dept='OSD运营支持部')
        try:
            self.balanced(v._build_a5_rows('耐数电子','2050',ws,total,2026,2),112)
            self.assertEqual(v._build_a6_rows('耐数电子','2050',ws,2026,2,timesheet()),[])
            self.assertEqual(v._build_a10_rows('耐数电子','2050',2026,2,expense(dept='OSD运营支持部'),timesheet()),[])
        finally:wb.close()


if __name__=='__main__':unittest.main()
