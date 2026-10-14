from decimal import Decimal
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '工作区域'))

from payroll_tool.options import RunOptions, build_batch_layout_from_options  # noqa: E402
from payroll_tool.pipeline import execute_payroll_run  # noqa: E402
from payroll_tool.precheck import run_startup_precheck  # noqa: E402
from payroll_tool.workbook_io import is_strict_xls, load_workbook  # noqa: E402


class PayrollIntegrationTests(unittest.TestCase):
    """Exercise a real, minimal Excel batch without using production data."""

    def _write_workbook(self, path, sheets):
        wb = Workbook()
        first_sheet = True
        for title, rows in sheets.items():
            ws = wb.active if first_sheet else wb.create_sheet()
            first_sheet = False
            ws.title = title
            for row in rows:
                ws.append(row)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)

    def _create_minimal_batch(self, root):
        base_dir = root / '工作区域'
        options = RunOptions(
            companies=('耐数电子',),
            vouchers=('A5', 'A6'),
            posting_year=2026,
            posting_month=3,
        )
        layout = build_batch_layout_from_options(str(base_dir), options)

        self._write_workbook(
            Path(layout.mapping_path),
            {
                '成本中心': [
                    ['序号', '2050部门', '2050成本中心', '', '', '', '2060部门', '2060成本中心'],
                    [1, '研发部', '20501001', '', '', '', '', ''],
                ]
            },
        )
        self._write_workbook(
            Path(layout.timesheet_path),
            {
                '工时汇总': [
                    ['序号', '年份', '月份', '', '部门', '公司', '项目', '', '', '类型', '工时', '', '内部订单'],
                    [1, 2026, 3, '', '研发部', '耐数电子', '项目A', '', '', '研发', 8, '', ''],
                ],
                'KOK3': [
                    ['内部订单', '', '', '', '', '', '', '', '订单描述', '', '', '公司代码'],
                    ['9201001', '', '', '', '', '', '', '', '项目A', '', '', 2050],
                ],
            },
        )
        raw_path = Path(layout.raw_dir) / '人力成本研发项目分摊202602 - to财务-原始.xlsx'
        self._write_workbook(
            raw_path,
            {
                '工资单': [
                    ['工资所属月份'],
                    [
                        '公司', '部门', '项目', '类别', '税前工资', '养老公司', '失业公司', '工伤公司', '医疗公司', '公积金公司',
                        '养老个人', '失业个人', '医疗个人', '公积金个人', '个税', '其他扣款', '实发金额', '备注',
                    ],
                    ['耐数电子', '研发部', '项目A', '薪资', 100, 0, 0, 0, 0, 0, 0, 0, 0, 0, 10, 0, 90, ''],
                    ['总计', '', '', '', 100, 0, 0, 0, 0, 0, 0, 0, 0, 0, 10, 0, 90, ''],
                ]
            },
        )

        # The app loads templates through BytesIO, so a workbook may retain the
        # historical .XLS filename while still being valid OOXML test content.
        template_path = root / '模版性文件' / '2060总账凭证导入-实际202602薪酬.XLS'
        self._write_workbook(template_path, {'凭证': [['组号']]})
        return base_dir, layout, raw_path, options

    def test_precheck_and_run_generate_balanced_accrual_vouchers(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            base_dir, layout, raw_path, options = self._create_minimal_batch(Path(temp_dir))

            precheck = run_startup_precheck(str(base_dir), layout.mapping_path, layout.bank_dir, options)
            self.assertTrue(precheck['can_run'], precheck.get('summary'))
            self.assertEqual(precheck['company_checks']['耐数电子']['voucher_status']['A5']['status'], 'ok')
            self.assertEqual(precheck['company_checks']['耐数电子']['voucher_status']['A6']['status'], 'ok')

            result = execute_payroll_run(
                str(raw_path), str(base_dir), layout.mapping_path, layout.bank_dir, lambda *_: None, run_options=options
            )
            voucher_path = Path(result['voucher_paths']['耐数电子（计提）'])
            self.assertTrue(voucher_path.is_file())
            self.assertTrue(result['voucher_validation_summary']['耐数电子']['a5_balanced'])
            self.assertTrue(result['voucher_validation_summary']['耐数电子']['a6_balanced'])

            self.assertTrue(is_strict_xls(voucher_path))
            wb = load_workbook(voucher_path, data_only=True)
            try:
                totals = {}
                for row in wb.active.iter_rows(min_row=2, max_col=13, values_only=True):
                    group, posting_key, amount = row[0], row[6], row[12]
                    if not group:
                        continue
                    debit, credit = totals.setdefault(group, [Decimal('0.00'), Decimal('0.00')])
                    if str(posting_key) == '40':
                        debit += Decimal(str(amount))
                    elif str(posting_key) == '50':
                        credit += Decimal(str(amount))
                    totals[group] = [debit, credit]
                    self.assertIsInstance(amount, str)
                self.assertEqual(set(totals), {'A5', 'A6'})
                for group, (debit, credit) in totals.items():
                    self.assertEqual(debit, credit, f'{group} 借贷不平：{debit} != {credit}')
                    self.assertEqual(debit, Decimal('100.00'))
            finally:
                wb.close()


if __name__ == '__main__':
    unittest.main()
