import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from decimal import Decimal
from openpyxl import Workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / '工作区域'))
from payroll_tool.publication import (publish_files, recover_pending, workspace_lock,
                                     merge_voucher, write_json, digest)
from payroll_tool.options import RunOptions
from payroll_tool.core import _make_voucher_row
from payroll_tool.workbook_io import load_workbook


def write_voucher(path, groups):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.append(['组号'])
    from datetime import date
    for group, amount in groups.items():
        for key in [40, 50]:
            row = _make_voucher_row(group,'2050',date(2026,3,31),'6601010001',amount,'text',key)
            ws.append([row.get(chr(65+i)) for i in range(17)])
    wb.save(path); wb.close()


class PublicationTests(unittest.TestCase):
    def test_partial_merge_preserves_unselected_and_removes_empty_selection(self):
        with tempfile.TemporaryDirectory() as tmp:
            base=Path(tmp); rel=Path('2050_耐数电子/03_计提凭证/test.XLS')
            old=base/'old'/rel; new=base/'stage'/rel
            write_voucher(old,{'A5':100,'A6':100}); write_voucher(new,{'A5':120})
            options=RunOptions(('耐数电子',),('A5',),2026,3)
            merge_voucher(new,old,options,'run1')
            wb=load_workbook(new)
            try:
                rows=list(wb.active.iter_rows(min_row=2,values_only=True))
                self.assertEqual([(r[0],r[12]) for r in rows],[('A5','120.00')]*2+[('A6','100.00')]*2)
            finally: wb.close()
            write_voucher(new,{})
            merge_voucher(new,old,options,'run2')
            wb=load_workbook(new)
            try: self.assertEqual({r[0] for r in wb.active.iter_rows(min_row=2,values_only=True)}, {'A6'})
            finally: wb.close()

    def test_corrupted_output_rejected(self):
        for delta in [Decimal('.01'),Decimal('1')]:
            with self.subTest(delta=delta), tempfile.TemporaryDirectory() as tmp:
                path=Path(tmp)/'2050_耐数电子/03_计提凭证/test.XLS'
                write_voucher(path,{'A5':100})
                wb=load_workbook(path); wb.active['M2']=f'{100+delta:.2f}'; wb.save(path); wb.close()
                with self.assertRaisesRegex(ValueError,'借贷不平'):
                    merge_voucher(path,path.parent/'missing.XLS',RunOptions(('耐数电子',),('A5',),2026,3),'run')

    def test_wrong_month_blocks_merge(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'2050_耐数电子/03_计提凭证/test.XLS'
            write_voucher(path,{'A5':100})
            with self.assertRaisesRegex(ValueError,'月份'):
                merge_voucher(path,path.parent/'missing.XLS',RunOptions(('耐数电子',),('A5',),2026,4),'run')

    def test_a3_merge_keeps_actual_groups_and_blocks_cross_account_difference(self):
        from test_business_regression import payroll,bank,timesheet
        from payroll_tool import vouchers as v
        with tempfile.TemporaryDirectory() as tmp:
            rel=Path('2050_耐数电子/02_实际凭证/test.XLS')
            old=Path(tmp)/'old'/rel; new=Path(tmp)/'stage'/rel
            wb,ws,total=payroll();b=bank('2050');t=timesheet()
            try:
                a3=v._build_a3_rows('耐数电子','2050',ws,b,2026,2,t)
                all_rows=(v._build_a1_rows('耐数电子','2050',ws,total,b,2026,2)+
                          v._build_a2_rows('耐数电子','2050',b,2026,2)+a3+
                          v._build_a4_rows('耐数电子','2050',ws,total,2026,2,t))
            finally:wb.close()
            def save(path,rows):
                path.parent.mkdir(parents=True,exist_ok=True)
                out=Workbook();out.active.append(['组号'])
                for row in rows:out.active.append([row.get(chr(65+i)) for i in range(17)])
                out.save(path);out.close()
            save(old,all_rows);save(new,a3)
            options=RunOptions(('耐数电子',),('A3',),2026,3)
            merge_voucher(new,old,options,'new')
            out=load_workbook(new)
            try:self.assertEqual({r[0] for r in out.active.iter_rows(min_row=2,values_only=True)},{'A1','A2','A3','A4'})
            finally:out.close()
            changed=[dict(r) for r in a3]
            for r in changed:
                if r['I']=='2211010001' or (r['G']==50 and r['M']=='87.00'):r['M']='88.00'
            save(new,changed)
            with self.assertRaisesRegex(ValueError,'对冲异常'):merge_voucher(new,old,options,'bad')

    def test_failed_second_publish_restores_first(self):
        with tempfile.TemporaryDirectory() as tmp:
            base=Path(tmp); history=base/'04_归档留痕/b/运行历史/r'; stage=history/'staging'
            a=stage/'03_运行输出/a.xlsx'; b=stage/'03_运行输出/b.xlsx'
            write_voucher(a,{'A5':120}); write_voucher(b,{'A5':120})
            target=base/'03_运行输出/a.xlsx'; write_voucher(target,{'A5':100}); before=digest(target)
            import os
            original=os.replace
            def fail(source,dest):
                if Path(dest)==base/'03_运行输出/b.xlsx': raise PermissionError('模拟文件占用')
                return original(source,dest)
            with patch('payroll_tool.publication.os.replace',side_effect=fail):
                with self.assertRaises(PermissionError): publish_files(base,history,stage,[a,b],'required')
            self.assertEqual(digest(target),before)
            self.assertFalse((base/'03_运行输出/b.xlsx').exists())
            recover_pending(base)

    def test_interrupted_publish_recovers_and_lock_is_exclusive(self):
        with tempfile.TemporaryDirectory() as tmp:
            base=Path(tmp); history=base/'04_归档留痕/b/运行历史/r'
            target=base/'03_运行输出/a.xlsx'; backup=history/'previous/a.xlsx'
            write_voucher(backup,{'A5':100}); write_voucher(target,{'A5':120})
            write_json(history/'publish-required.json',{'status':'publishing','entries':[
                dict(target=str(target),backup=str(backup),existed=True,old_hash=digest(backup))]})
            with workspace_lock(base):
                with self.assertRaises(RuntimeError):
                    with workspace_lock(base): pass
                recover_pending(base)
            self.assertEqual(digest(target),digest(backup))


if __name__ == '__main__': unittest.main()
