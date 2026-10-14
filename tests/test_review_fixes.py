"""Regression cases for the code review, without production data."""
import sys
import tempfile
import unittest
from pathlib import Path
from decimal import Decimal
from unittest.mock import patch
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / '工作区域'))
from payroll_tool.bank_identity import TransactionIndex, BankIdentityError
from payroll_tool.contexts import _match_timesheet_internal_orders, _load_timesheet_match_context
from payroll_tool.vouchers import _build_a9_rows, _build_a10_rows
from payroll_tool.validation import _validate_voucher_groups
from payroll_tool.workbook_io import load_workbook, read_voucher_values, convert_to_strict_xls, excel_scope
from test_bank_identity import record
from test_business_regression import expense, timesheet


class ReviewFixTests(unittest.TestCase):
    def reversal(self, amount='100', usage='支付平台退票：收款人名称有误'):
        result = record(tid='1', amount='0', row=3)
        result.update(incoming_amt=Decimal(amount), usage=usage)
        return result

    def test_return_keeps_both_legs_and_deduplicates_copies(self):
        for reverse in (False, True):
            index = TransactionIndex()
            legs = [record(tid='1'), self.reversal()]
            if reverse:
                legs.reverse()
            for leg in legs:
                self.assertTrue(index.add(dict(leg)))
            for leg in legs:
                self.assertFalse(index.add(dict(leg, file='copy.xlsx')))
            self.assertEqual(len(index.records), 2)
            self.assertEqual(sum(r['incoming_amt'] - r['outgoing_amt'] for r in index.records), 0)
            self.assertEqual(index.deduped, 2)

    def test_unexplained_or_unequal_reverse_still_blocks(self):
        for reverse in (self.reversal(usage='转账'), self.reversal(amount='101'),
                        dict(self.reversal(), counterparty_account='different')):
            index = TransactionIndex()
            index.add(record(tid='1'))
            with self.assertRaises(BankIdentityError):
                index.add(reverse)

    def test_amount_conflict_after_return_still_blocks(self):
        index = TransactionIndex()
        index.add(record(tid='1'))
        index.add(self.reversal())
        with self.assertRaises(BankIdentityError):
            index.add(record(tid='1', amount='101'))

    def test_explicit_payment_platform_return_is_preserved(self):
        index = TransactionIndex()
        index.add(record(tid='1'))
        self.assertTrue(index.add(self.reversal(usage='支付平台退回：收付款人不为同名户')))

    def test_generation_uses_same_kok3_match_as_precheck(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'timesheet.xlsx'
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = '工时汇总'
            sheet.append(['header'] * 13)
            sheet.append([None, 2026, 9, None, '研发部', '耐数电子', 'PROJECT',
                          None, None, '自研', 10, None, '9209999'])
            kok3 = workbook.create_sheet('KOK3')
            kok3.append(['header'] * 12)
            kok3.append(['9201001', None, None, None, None, None, None, None,
                         'PROJECT', None, None, 2050])
            workbook.save(path)
            workbook.close()
            before = path.read_bytes()
            matched, _ = _match_timesheet_internal_orders(path, save_changes=False)
            context = _load_timesheet_match_context(path)
            self.assertEqual(context['project_index'], matched)
            self.assertEqual(context['allocation_index'][(2026, 9, '耐数电子', '研发部')],
                             {'9201001': Decimal('10')})
            self.assertEqual(path.read_bytes(), before)

    def test_alternate_statement_requires_full_payment_count_match(self):
        from payroll_tool.compensation_report import _load_bank_records
        for standard_count in (1, 2):
            with self.subTest(standard_count=standard_count), tempfile.TemporaryDirectory() as temp:
                root = Path(temp)
                workbook = Workbook()
                sheet = workbook.active
                sheet.cell(7, 2, '人民币')
                for col, name in ((1, '账号'), (4, '交易日'), (8, '借方金额')):
                    sheet.cell(13, col, name)
                for row in (14, 15):
                    for col, value in {1: 'account', 4: '2025-12-05', 8: 100,
                                       11: '代发工资', 20: '员工', 21: 'counterparty'}.items():
                        sheet.cell(row, col, value)
                workbook.save(root / '2050招行对帐单.xlsx')
                workbook.close()
                workbook = Workbook()
                sheet = workbook.active
                sheet.append(['header'] * 40)
                sheet.cell(1, 32, '交易流水号')
                for row in range(2, 2 + standard_count):
                    for col, value in {2: '2050', 3: 'account', 4: '2025-12-05',
                                       5: 'CNY', 13: '员工', 14: 0, 15: 100,
                                       16: '代发工资 代发工资', 20: 'counterparty',
                                       32: str(row)}.items():
                        sheet.cell(row, col, value)
                workbook.save(root / '银行流水-205012.xlsx')
                workbook.close()
                if standard_count == 1:
                    with self.assertRaises(BankIdentityError):
                        _load_bank_records(root)
                else:
                    records, stats = _load_bank_records(root)
                    self.assertEqual(len(records[('2050', 2025, 12)]), 2)
                    self.assertEqual(stats['deduped_row_count'], 2)
                    self.assertTrue(any('全月付款及次数一致' in w for w in stats['warnings']))

    def test_a9_small_amounts_both_directions_and_repeated_orders(self):
        for amount in ('0.01', '0.02', '-0.01', '-0.02'):
            context = {
                'allocation_rows': [{'order': str(9201001 + i), 'ratio_base': Decimal('1')}
                                    for i in range(4)],
                'total_ratio_base': Decimal('4'),
                'expense_rows': [{'source_account': '6601010003', 'debit_gl': '5001010011',
                                  'amount': Decimal(amount)}],
            }
            rows = _build_a9_rows('2050', 2026, 8, context)
            _validate_voucher_groups(rows)
            self.assertTrue(all(Decimal(row['M']) > 0 for row in rows))
            for side in (40, 50):
                self.assertEqual(sum(Decimal(r['M']) for r in rows if r['G'] == side),
                                 abs(Decimal(amount)))

    def test_a9_normal_rounding_is_preserved_with_repeated_order(self):
        context = {
            'allocation_rows': [{'order': order, 'ratio_base': Decimal('1')}
                                for order in ('9201001', '9201002', '9201002')],
            'total_ratio_base': Decimal('3'),
            'expense_rows': [{'source_account': '6601010003', 'debit_gl': '5001010011',
                              'amount': Decimal('1')}],
        }
        rows = _build_a9_rows('2050', 2026, 8, context)
        self.assertEqual([(r['K'], r['M']) for r in rows if r['G'] == 40],
                         [('9201001', '0.33'), ('9201002', '0.67')])

    def test_a10_is_not_skipped_by_a5_only_rule(self):
        rows = _build_a10_rows('耐数信息', '2060', 2026, 8,
                              expense('耐数信息', 'ATS测试解决方案', '20601001'),
                              timesheet('耐数信息', 'ATS测试解决方案', 9))
        _validate_voucher_groups(rows)
        self.assertEqual(len(rows), 2)
        self.assertEqual([r['M'] for r in rows], ['10.00', '10.00'])

    def test_bytes_path_is_supported(self):
        import os
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'input.xlsx'
            workbook = Workbook()
            workbook.active['A1'] = 'test'
            workbook.save(path)
            workbook.close()
            workbook = load_workbook(os.fsencode(path))
            try:
                self.assertEqual(workbook.active['A1'].value, 'test')
            finally:
                workbook.close()

    def test_strict_xls_values_are_checked_without_second_excel_start(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'voucher.XLS'
            workbook = Workbook()
            workbook.active.append(['header'] * 17)
            workbook.active.append(['A9', '2050', 'SA', '20260930', '20260930', 'CNY',
                                    40, None, '5001010011', '', '9201001', '', '0.02',
                                    'test', None, None, None])
            workbook.save(path)
            workbook.close()
            import win32com.client
            with patch('win32com.client.DispatchEx', wraps=win32com.client.DispatchEx) as start, excel_scope():
                values = convert_to_strict_xls(path, read_values=True)
                self.assertEqual(read_voucher_values(path), values)
                self.assertEqual(start.call_count, 1)
            self.assertEqual(values[0][6], 40)
            self.assertEqual(values[0][12], '0.02')
            self.assertEqual(len(values[0]), 17)


if __name__ == '__main__':
    unittest.main()
