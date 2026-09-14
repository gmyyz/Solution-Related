import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'工作区域'))
import test_payroll_integration as integration
from payroll_tool.pipeline import execute_payroll_run
from payroll_tool.publication import digest
from payroll_tool import core


class RunSafetyTests(unittest.TestCase):
    def setup_batch(self,root):
        return integration.PayrollIntegrationTests()._create_minimal_batch(root)

    def run_batch(self,batch):
        base,layout,raw,options=batch
        return execute_payroll_run(str(raw),str(base),layout.mapping_path,layout.bank_dir,lambda *args:None,run_options=options)

    def test_second_run_retains_history_and_source_hash(self):
        with tempfile.TemporaryDirectory() as tmp:
            batch=self.setup_batch(Path(tmp)); before=digest(batch[2])
            first=self.run_batch(batch); second=self.run_batch(batch)
            self.assertNotEqual(first['run_id'],second['run_id'])
            self.assertTrue(Path(first['history_path'],'run.json').exists())
            self.assertEqual(digest(batch[2]),before)
            self.assertTrue(list(Path(second['history_path']).rglob('previous-required/**/*.XLS')))

    def test_later_generation_failure_does_not_replace_existing(self):
        with tempfile.TemporaryDirectory() as tmp:
            batch=self.setup_batch(Path(tmp)); first=self.run_batch(batch)
            target=Path(next(iter(first['voucher_paths'].values()))); before=digest(target)
            with patch('payroll_tool.core._generate_voucher_files',side_effect=ValueError('第二家公司失败')):
                with self.assertRaises(ValueError):self.run_batch(batch)
            self.assertEqual(digest(target),before)

    def test_input_change_blocks_publication(self):
        with tempfile.TemporaryDirectory() as tmp:
            batch=self.setup_batch(Path(tmp)); first=self.run_batch(batch)
            target=Path(next(iter(first['voucher_paths'].values()))); before=digest(target)
            original=core.fill_first_sheet_ab
            def changing(*args,**kwargs):
                result=original(*args,**kwargs)
                # Test fixture only: change input after generation, before publication.
                from openpyxl import load_workbook
                wb=load_workbook(batch[2]);wb.active['A1']='changed';wb.save(batch[2]);wb.close()
                return result
            with patch('payroll_tool.pipeline.fill_first_sheet_ab',side_effect=changing):
                with self.assertRaisesRegex(ValueError,'输入已变化'):self.run_batch(batch)
            self.assertEqual(digest(target),before)

    def test_optional_report_failure_still_delivers_vouchers(self):
        with tempfile.TemporaryDirectory() as tmp:
            batch=self.setup_batch(Path(tmp))
            with patch('payroll_tool.core.generate_compensation_report',side_effect=ValueError('历史资料需核对')):
                result=self.run_batch(batch)
            self.assertEqual(result['status'],'success_with_warnings')
            self.assertTrue(result['optional_errors'])
            self.assertTrue(all(Path(p).is_file() for p in result['voucher_paths'].values()))
            self.assertNotIn('compensation_report',result['artifact_paths'])

    def test_equal_debit_credit_corruption_is_still_detected(self):
        with tempfile.TemporaryDirectory() as tmp:
            batch=self.setup_batch(Path(tmp))
            from payroll_tool import voucher_output
            original=voucher_output._write_voucher_rows_to_template
            def corrupt(ws,rows):
                original(ws,rows)
                # Both sides remain balanced, but no longer match generated rows.
                for row in range(2,ws.max_row+1):
                    if ws.cell(row,1).value=='A5':
                        from decimal import Decimal
                        ws.cell(row,13).value=f'{Decimal(ws.cell(row,13).value)+1:.2f}'
            with patch.object(voucher_output,'_write_voucher_rows_to_template',side_effect=corrupt):
                with self.assertRaisesRegex(ValueError,'生成记录不一致'):self.run_batch(batch)


if __name__=='__main__':unittest.main()
