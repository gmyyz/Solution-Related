"""tax_report: extracted payroll responsibilities; business rules preserved."""
from copy import copy
from decimal import Decimal, ROUND_HALF_UP, ROUND_DOWN
from openpyxl import Workbook
from .workbook_io import load_workbook, workbook_scope
from .banking import (
    bank_scope,
    _get_treasury_records,
    _load_bank_records,
    _load_bank_records_with_stats,
    _match_bank_record_combination,
    _match_salary_combination,
    _match_treasury_combination,
    _normalize_bank_date_key,
    _summarize_bank_records,
)
from .workbook_utils import (
    _copy_cell_style,
    _find_payroll_header_row,
    _find_total_row,
    _format_code,
    _is_blank,
    _is_blank_project,
    _is_total_row,
    _mark_range_red,
    _mark_row_red,
    _normalize_text,
    _recalculate_total_row,
    _to_decimal,
    _to_money,
    _to_text_money,
)
from .domain import (
    COMPANY_NAME_TO_CODE,
    _amount_to_ten_thousand_yuan,
    _get_bonus_tax_adjustment,
    _shift_month,
)
from .paths import (
    _build_key_tax_source_output_path,
)


def _describe_bank_records(records):
    if not records:
        return ''
    return '；'.join(
        f'{record["file"]} 第{record["row_idx"]}行 {record["outgoing_amt"]}'
        for record in records
    )


def _append_key_tax_source_row(rows, company, company_code, processing_label, payroll_label, category, item, amount, source, note):
    amount = _to_money(amount)
    rows.append(
        {
            '处理月份': processing_label,
            '工资所属月份': payroll_label,
            '公司代码': company_code,
            '公司': company,
            '类别': category,
            '项目': item,
            '金额（元）': amount,
            '金额（万元）': _amount_to_ten_thousand_yuan(amount),
            '取数口径': source,
            '备注': note,
        }
    )


@workbook_scope
def _build_company_key_tax_source_rows(company, payroll_path, bank_records, payroll_year, payroll_month):
    company_code = COMPANY_NAME_TO_CODE[company]
    processing_year, processing_month = _shift_month(payroll_year, payroll_month, 1)
    processing_label = f'{processing_year}年{processing_month}月'
    payroll_label = f'{payroll_year}年{payroll_month}月'

    wb = load_workbook(payroll_path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    total_row_idx = _find_total_row(ws)
    if total_row_idx is None:
        raise ValueError(f'{company} 工资整理文件中未找到总计行，无法生成重点税源采集信息')

    company_pension = _to_money(ws.cell(total_row_idx, 6).value)
    company_unemployment = _to_money(ws.cell(total_row_idx, 7).value)
    company_injury = _to_money(ws.cell(total_row_idx, 8).value)
    company_medical = _to_money(ws.cell(total_row_idx, 9).value)
    social_total = sum(
        (_to_money(ws.cell(total_row_idx, col_idx).value) for col_idx in (6, 7, 8, 9, 11, 12, 13)),
        Decimal('0.00'),
    )

    salary_tax = Decimal('0.00')
    labor_tax = Decimal('0.00')
    data_start_row = _find_payroll_header_row(ws) + 1
    for row_idx, ws_cells in enumerate(ws.iter_rows(min_row=data_start_row), data_start_row):
        if _is_total_row(ws_cells[(1) - 1].value):
            break
        item_type = _normalize_text(ws_cells[(4) - 1].value)
        if item_type == '薪资':
            salary_tax += _to_decimal(ws_cells[(15) - 1].value)
        elif item_type == '劳务费':
            labor_tax += _to_decimal(ws_cells[(15) - 1].value)
    salary_tax = salary_tax.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    labor_tax = labor_tax.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    bonus_tax = _get_bonus_tax_adjustment(company_code, processing_year, processing_month)
    salary_tax_with_bonus = (salary_tax + bonus_tax).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    tax_total = (salary_tax_with_bonus + labor_tax).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

    month_records = bank_records.get((company_code, processing_year, processing_month), [])
    social_match = _match_treasury_combination(month_records, social_total, '社保国库流水')
    social_note = social_match['note']
    if social_match['matched_records']:
        social_note = f'{social_note}；来源：{_describe_bank_records(social_match["matched_records"])}'

    fund_records = [record for record in month_records if '北京住房公积金管理中心' in record['payment_name']]
    fund_total = sum((record['outgoing_amt'] for record in fund_records), Decimal('0.00'))
    company_fund = (fund_total / Decimal('2')).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    fund_note = f'银行流水公积金实付合计 {fund_total}，按现有凭证口径公司承担 50%'
    if fund_records:
        fund_note = f'{fund_note}；来源：{_describe_bank_records(fund_records)}'

    tax_match = _match_treasury_combination(month_records, tax_total, '个税国库流水')
    tax_note = tax_match['note']
    if tax_match['matched_records']:
        tax_note = f'{tax_note}；来源：{_describe_bank_records(tax_match["matched_records"])}'

    rows = []
    _append_key_tax_source_row(
        rows,
        company,
        company_code,
        processing_label,
        payroll_label,
        '社保',
        '养老保险（公司承担）',
        company_pension,
        '工资整理文件总计行 F列，公司承担养老；社保国库流水用于实缴匹配',
        social_note,
    )
    _append_key_tax_source_row(
        rows,
        company,
        company_code,
        processing_label,
        payroll_label,
        '社保',
        '医疗保险（公司承担）',
        company_medical,
        '工资整理文件总计行 I列，公司承担医疗；社保国库流水用于实缴匹配',
        social_note,
    )
    _append_key_tax_source_row(
        rows,
        company,
        company_code,
        processing_label,
        payroll_label,
        '社保',
        '失业保险（公司承担）',
        company_unemployment,
        '工资整理文件总计行 G列，公司承担失业；社保国库流水用于实缴匹配',
        social_note,
    )
    _append_key_tax_source_row(
        rows,
        company,
        company_code,
        processing_label,
        payroll_label,
        '社保',
        '工伤保险（公司承担）',
        company_injury,
        '工资整理文件总计行 H列，公司承担工伤；社保国库流水用于实缴匹配',
        social_note,
    )
    _append_key_tax_source_row(
        rows,
        company,
        company_code,
        processing_label,
        payroll_label,
        '公积金',
        '住房公积金（公司承担）',
        company_fund,
        '银行流水付款名称含“北京住房公积金管理中心”的实付金额，按公司承担 50% 拆分',
        fund_note,
    )
    _append_key_tax_source_row(
        rows,
        company,
        company_code,
        processing_label,
        payroll_label,
        '个人所得税',
        '个人所得税-工资薪金',
        salary_tax_with_bonus,
        '工资整理文件薪资行 O列个税；若存在本月已确认年终奖个税调整，则并入工资薪金口径',
        f'薪资个税 {salary_tax}，年终奖个税调整 {bonus_tax}；{tax_note}',
    )
    _append_key_tax_source_row(
        rows,
        company,
        company_code,
        processing_label,
        payroll_label,
        '个人所得税',
        '个人所得税-劳务',
        labor_tax,
        '工资整理文件劳务费行 O列个税',
        tax_note,
    )
    _append_key_tax_source_row(
        rows,
        company,
        company_code,
        processing_label,
        payroll_label,
        '个人所得税',
        '个人所得税合计',
        tax_total,
        '工资薪金个税与劳务个税合计，并与个税国库流水匹配',
        tax_note,
    )
    wb.close()
    return rows


def _write_key_tax_source_workbook(report_path, rows, processing_year, processing_month):
    wb = Workbook()
    ws = wb.active
    ws.title = '重点税源采集信息'
    headers = ['处理月份', '工资所属月份', '公司代码', '公司', '类别', '项目', '金额（元）', '金额（万元）', '取数口径', '备注']
    ws.append(headers)
    for row in rows:
        ws.append([row[header] for header in headers])

    widths = [14, 14, 10, 12, 12, 24, 14, 14, 46, 70]
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, col_idx).column_letter].width = width
    ws.freeze_panes = 'A2'
    for cell in ws[1]:
        font = copy(cell.font)
        font.bold = True
        cell.font = font
    for row_idx in range(2, ws.max_row + 1):
        ws.cell(row_idx, 7).number_format = '#,##0.00'
        ws.cell(row_idx, 8).number_format = '#,##0.0000'

    note_ws = wb.create_sheet('口径说明')
    note_ws.append(['项目', '说明'])
    note_ws.append(['输出位置', report_path])
    note_ws.append(['统计月份', f'{processing_year}年{processing_month}月银行流水'])
    note_ws.append(['社保公司承担', '养老、医疗、失业、工伤取工资整理文件公司承担列，并用对应社保国库流水匹配实缴总额。'])
    note_ws.append(['公积金公司承担', '取北京住房公积金管理中心银行实付金额的 50%，沿用现有 A2 凭证拆分口径。'])
    note_ws.append(['个人所得税', '工资薪金个税包含薪资行个税及本月已确认年终奖个税调整；劳务个税单独列示。'])
    note_ws.column_dimensions['A'].width = 18
    note_ws.column_dimensions['B'].width = 100
    wb.save(report_path)


def _write_key_tax_source_report(base_dir, output_paths, bank_records, payroll_year, payroll_month):
    processing_year, processing_month = _shift_month(payroll_year, payroll_month, 1)
    report_paths = {}
    for company, payroll_path in output_paths.items():
        rows = _build_company_key_tax_source_rows(company, payroll_path, bank_records, payroll_year, payroll_month)
        report_path = _build_key_tax_source_output_path(base_dir, company, processing_year, processing_month)
        _write_key_tax_source_workbook(report_path, rows, processing_year, processing_month)
        report_paths[company] = report_path
    return report_paths
