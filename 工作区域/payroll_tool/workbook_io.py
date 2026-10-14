"""Owned workbook resources; file handles are released even if parsing fails."""
from contextvars import ContextVar
from contextlib import contextmanager
from functools import wraps
from io import BytesIO
import os
from os import PathLike
from pathlib import Path
import shutil
import tempfile
from openpyxl import load_workbook as _load

_owned = ContextVar('payroll_owned_workbooks', default=None)
_excel_session = ContextVar('payroll_excel_session', default=None)
_OLE_SIGNATURE = bytes.fromhex('D0CF11E0A1B11AE1')
_OOXML_SIGNATURE = b'PK\x03\x04'


def _file_signature(path):
    with open(path, 'rb') as stream:
        return stream.read(8)


def is_strict_xls(path):
    """Return whether a workbook is a real Excel 97-2003 OLE/BIFF file."""
    return _file_signature(path) == _OLE_SIGNATURE


def _close_excel_session(session):
    try:
        if session.get('excel') is not None:
            session['excel'].Quit()
    finally:
        if session.get('pythoncom') is not None:
            session['pythoncom'].CoUninitialize()


@contextmanager
def excel_scope():
    """Reuse one private Excel process on this thread for one batch only."""
    if _excel_session.get() is not None:
        yield
        return
    session = {}
    token = _excel_session.set(session)
    try:
        yield
    finally:
        try:
            _close_excel_session(session)
        finally:
            _excel_session.reset(token)


def _convert_with_excel(source, destination=None, file_format=None, read_values=False):
    """Use desktop Excel for a real format conversion, not an extension rename."""
    if os.name != 'nt':
        raise RuntimeError('生成严格 XLS 需要在 Windows 上安装 Microsoft Excel')
    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:
        raise RuntimeError('生成严格 XLS 需要安装 pywin32，请重新执行依赖安装') from exc

    source = Path(source).resolve()
    destination = Path(destination).resolve() if destination is not None else None
    session = _excel_session.get()
    owns_session = session is None
    if owns_session:
        session = {}
    workbook = None
    try:
        if session.get('excel') is None:
            pythoncom.CoInitialize()
            session['pythoncom'] = pythoncom
            excel = win32com.client.DispatchEx('Excel.Application')
            session['excel'] = excel
            excel.Visible = False
            excel.DisplayAlerts = False
        else:
            excel = session['excel']
        workbook = excel.Workbooks.Open(
            str(source),
            UpdateLinks=0,
            ReadOnly=True,
            IgnoreReadOnlyRecommended=True,
        )
        if destination is not None:
            workbook.CheckCompatibility = False
            workbook.SaveAs(str(destination), FileFormat=file_format)
            if read_values:
                workbook.Close(SaveChanges=False)
                workbook = None
                workbook = excel.Workbooks.Open(str(destination), UpdateLinks=0, ReadOnly=True)
        if read_values:
            sheet = workbook.Worksheets(1)
            last_row = sheet.UsedRange.Row + sheet.UsedRange.Rows.Count - 1
            if last_row < 2:
                return []
            values = sheet.Range(f'A2:Q{last_row}').Value
            def normalized(value):
                if isinstance(value, float) and value.is_integer():
                    return int(value)
                return value
            return [tuple(normalized(v) for v in row) for row in values
                    if any(v is not None and v != '' for v in row)]
    except Exception as exc:
        raise RuntimeError(f'Excel 格式转换失败：{source.name}') from exc
    finally:
        try:
            if workbook is not None:
                workbook.Close(SaveChanges=False)
        finally:
            if owns_session:
                _close_excel_session(session)


def convert_to_strict_xls(path, read_values=False):
    """Atomically replace an OOXML workbook with a real Excel 97-2003 XLS."""
    path = Path(path)
    if is_strict_xls(path):
        return read_voucher_values(path) if read_values else None
    if _file_signature(path)[:4] != _OOXML_SIGNATURE:
        raise ValueError(f'无法识别待转换的 Excel 文件格式：{path}')

    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.xls-convert-', dir=path.parent) as temp_dir:
        temp_dir = Path(temp_dir)
        source = temp_dir / 'source.xlsx'
        converted = temp_dir / 'converted.xls'
        shutil.copy2(path, source)
        values = _convert_with_excel(source, converted, 56, read_values=read_values)  # xlExcel8
        if not converted.is_file() or not is_strict_xls(converted):
            raise RuntimeError(f'Excel 未生成严格 XLS 文件：{path.name}')
        os.replace(converted, path)
        return values


def load_workbook(filename, *args, **kwargs):
    # Detach the source handle before parsing. This also supports the historical
    # OOXML templates carrying an .XLS extension.
    if isinstance(filename, (str, bytes, PathLike)):
        path = Path(os.fsdecode(filename)) if isinstance(filename, bytes) else Path(filename)
        if is_strict_xls(path):
            with tempfile.TemporaryDirectory(prefix='.xls-read-') as temp_dir:
                converted = Path(temp_dir) / 'workbook.xlsx'
                _convert_with_excel(path, converted, 51)  # xlOpenXMLWorkbook
                with open(converted, 'rb') as stream:
                    filename = BytesIO(stream.read())
        else:
            with open(path, 'rb') as stream:
                filename = BytesIO(stream.read())
    workbook = _load(filename, *args, **kwargs)
    resources = _owned.get()
    if resources is not None:
        resources.append(workbook)
    return workbook


def read_voucher_values(path):
    """Read persisted A:Q values, including real BIFF XLS files."""
    if not is_strict_xls(path):
        wb = load_workbook(path, data_only=True)
        try:
            return [row for row in wb.worksheets[0].iter_rows(min_row=2, max_col=17, values_only=True)
                    if any(value is not None and value != '' for value in row)]
        finally:
            wb.close()
    return _convert_with_excel(path, read_values=True)


def workbook_scope(function):
    @wraps(function)
    def wrapped(*args, **kwargs):
        resources = []
        token = _owned.set(resources)
        try:
            return function(*args, **kwargs)
        finally:
            try:
                for workbook in reversed(resources):
                    workbook.close()
            finally:
                _owned.reset(token)
    return wrapped
